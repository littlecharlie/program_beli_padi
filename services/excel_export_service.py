"""
Excel Export Service
Handles exporting various reports to Excel format
"""
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Dict, Any
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from sqlalchemy import func

from models.purchase_bill import PurchaseBill
from models.delivery_invoice import DeliveryInvoice
from models.farmer import Farmer
from models.rice_mill import RiceMill


class ExcelExportError(Exception):
    """Custom exception for Excel export errors"""
    pass


class ExcelExportService:
    """Service for exporting data to Excel files"""

    # Style definitions
    HEADER_FILL = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    TITLE_FONT = Font(bold=True, size=14)
    SUBTITLE_FONT = Font(size=10, italic=True)
    CURRENCY_FORMAT = '#,##0.00'
    NUMBER_FORMAT = '#,##0.00'
    DATE_FORMAT = 'DD/MM/YYYY HH:MM'
    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    @staticmethod
    def _create_workbook() -> Workbook:
        """Create a new workbook with default settings"""
        wb = Workbook()
        return wb

    @staticmethod
    def _format_header_row(ws, row: int, num_columns: int):
        """Apply header formatting to a row"""
        for col in range(1, num_columns + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = ExcelExportService.HEADER_FONT
            cell.fill = ExcelExportService.HEADER_FILL
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = ExcelExportService.THIN_BORDER

    @staticmethod
    def _format_currency_column(ws, col: int, start_row: int, end_row: int):
        """Format a column as currency"""
        for row in range(start_row, end_row + 1):
            cell = ws.cell(row=row, column=col)
            cell.number_format = ExcelExportService.CURRENCY_FORMAT
            cell.border = ExcelExportService.THIN_BORDER

    @staticmethod
    def _format_number_column(ws, col: int, start_row: int, end_row: int):
        """Format a column as number"""
        for row in range(start_row, end_row + 1):
            cell = ws.cell(row=row, column=col)
            cell.number_format = ExcelExportService.NUMBER_FORMAT
            cell.border = ExcelExportService.THIN_BORDER

    @staticmethod
    def _format_date_column(ws, col: int, start_row: int, end_row: int):
        """Format a column as date"""
        for row in range(start_row, end_row + 1):
            cell = ws.cell(row=row, column=col)
            cell.number_format = ExcelExportService.DATE_FORMAT
            cell.border = ExcelExportService.THIN_BORDER

    @staticmethod
    def _auto_adjust_column_width(ws):
        """Auto-adjust column widths based on content"""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass

            adjusted_width = min(max_length + 2, 50)  # Cap at 50 characters
            ws.column_dimensions[column_letter].width = adjusted_width

    @staticmethod
    def _add_title_section(ws, title: str, subtitle: str, start_row: int = 1) -> int:
        """Add title and subtitle section to worksheet"""
        # Title
        ws.cell(row=start_row, column=1, value=title)
        ws.cell(row=start_row, column=1).font = ExcelExportService.TITLE_FONT

        # Subtitle
        ws.cell(row=start_row + 1, column=1, value=subtitle)
        ws.cell(row=start_row + 1, column=1).font = ExcelExportService.SUBTITLE_FONT

        return start_row + 3  # Return the next available row

    @classmethod
    def export_purchase_bills(
        cls,
        db: Session,
        file_path: str,
        from_date: datetime,
        to_date: datetime,
        filter_id: Optional[int] = None
    ) -> str:
        """
        Export purchase bills to Excel

        Args:
            db: Database session
            file_path: Output file path
            from_date: Start date for filtering
            to_date: End date for filtering
            filter_id: Optional farmer ID to filter by

        Returns:
            Path to the generated Excel file
        """
        try:
            # Query purchase bills
            query = db.query(PurchaseBill).filter(
                PurchaseBill.bill_date >= from_date,
                PurchaseBill.bill_date <= to_date
            )

            if filter_id:
                query = query.filter(PurchaseBill.farmer_id == filter_id)

            bills = query.order_by(PurchaseBill.bill_date.desc()).all()

            # Create workbook
            wb = cls._create_workbook()
            ws = wb.active
            ws.title = "Purchase Bills"

            # Add title section
            date_range = f"{from_date.strftime('%d/%m/%Y')} - {to_date.strftime('%d/%m/%Y')}"
            next_row = cls._add_title_section(
                ws,
                "Purchase Bills Report",
                f"Date Range: {date_range} | Generated: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            )

            # Headers
            headers = [
                "Bill Number", "Date", "Farmer Name", "Farmer IC",
                "Truck Number", "Gross Weight (kg)", "Discount %",
                "Discount Weight (kg)", "Net Weight (kg)", "Price per 1000kg (RM)",
                "Total Payment (RM)", "Subsidy Estimate (RM)", "Status", "Created By"
            ]

            for col, header in enumerate(headers, 1):
                ws.cell(row=next_row, column=col, value=header)

            cls._format_header_row(ws, next_row, len(headers))

            # Data rows
            data_start_row = next_row + 1
            for row_idx, bill in enumerate(bills, data_start_row):
                ws.cell(row=row_idx, column=1, value=bill.bill_number)
                ws.cell(row=row_idx, column=2, value=bill.bill_date)
                ws.cell(row=row_idx, column=3, value=bill.farmer.name if bill.farmer else "Unknown")
                ws.cell(row=row_idx, column=4, value=bill.farmer.ic_number if bill.farmer else "N/A")
                ws.cell(row=row_idx, column=5, value=bill.truck.truck_number if bill.truck else "Unknown")
                ws.cell(row=row_idx, column=6, value=float(bill.gross_weight))
                ws.cell(row=row_idx, column=7, value=float(bill.total_discount_percent))
                ws.cell(row=row_idx, column=8, value=float(bill.discount_weight))
                ws.cell(row=row_idx, column=9, value=float(bill.net_weight))
                ws.cell(row=row_idx, column=10, value=float(bill.rice_price_per_1000kg))
                ws.cell(row=row_idx, column=11, value=float(bill.total_payment))
                ws.cell(row=row_idx, column=12, value=float(bill.subsidy_estimate))
                ws.cell(row=row_idx, column=13, value="Delivered" if bill.is_delivered else "Pending")
                ws.cell(row=row_idx, column=14, value=bill.created_by or "System")

            # Format columns
            data_end_row = data_start_row + len(bills) - 1
            if len(bills) > 0:
                cls._format_date_column(ws, 2, data_start_row, data_end_row)  # Date
                cls._format_number_column(ws, 6, data_start_row, data_end_row)  # Gross Weight
                cls._format_number_column(ws, 7, data_start_row, data_end_row)  # Discount %
                cls._format_number_column(ws, 8, data_start_row, data_end_row)  # Discount Weight
                cls._format_number_column(ws, 9, data_start_row, data_end_row)  # Net Weight
                cls._format_currency_column(ws, 10, data_start_row, data_end_row)  # Price
                cls._format_currency_column(ws, 11, data_start_row, data_end_row)  # Payment
                cls._format_currency_column(ws, 12, data_start_row, data_end_row)  # Subsidy

            # Add summary row
            summary_row = data_end_row + 2
            ws.cell(row=summary_row, column=1, value="TOTAL")
            ws.cell(row=summary_row, column=1).font = Font(bold=True)

            total_weight = sum(float(b.net_weight) for b in bills)
            total_payment = sum(float(b.total_payment) for b in bills)
            total_subsidy = sum(float(b.subsidy_estimate) for b in bills)

            ws.cell(row=summary_row, column=9, value=total_weight)
            ws.cell(row=summary_row, column=11, value=total_payment)
            ws.cell(row=summary_row, column=12, value=total_subsidy)

            ws.cell(row=summary_row, column=9).font = Font(bold=True)
            ws.cell(row=summary_row, column=11).font = Font(bold=True)
            ws.cell(row=summary_row, column=12).font = Font(bold=True)

            ws.cell(row=summary_row, column=9).number_format = cls.NUMBER_FORMAT
            ws.cell(row=summary_row, column=11).number_format = cls.CURRENCY_FORMAT
            ws.cell(row=summary_row, column=12).number_format = cls.CURRENCY_FORMAT

            # Auto-adjust column widths
            cls._auto_adjust_column_width(ws)

            # Save file
            wb.save(file_path)
            return file_path

        except Exception as e:
            raise ExcelExportError(f"Failed to export purchase bills: {str(e)}")

    @classmethod
    def export_delivery_invoices(
        cls,
        db: Session,
        file_path: str,
        from_date: datetime,
        to_date: datetime,
        filter_id: Optional[int] = None
    ) -> str:
        """
        Export delivery invoices to Excel

        Args:
            db: Database session
            file_path: Output file path
            from_date: Start date for filtering
            to_date: End date for filtering
            filter_id: Optional mill ID to filter by

        Returns:
            Path to the generated Excel file
        """
        try:
            # Query delivery invoices
            query = db.query(DeliveryInvoice).filter(
                DeliveryInvoice.invoice_date >= from_date,
                DeliveryInvoice.invoice_date <= to_date
            )

            if filter_id:
                query = query.filter(DeliveryInvoice.mill_id == filter_id)

            invoices = query.order_by(DeliveryInvoice.invoice_date.desc()).all()

            # Create workbook
            wb = cls._create_workbook()
            ws = wb.active
            ws.title = "Delivery Invoices"

            # Add title section
            date_range = f"{from_date.strftime('%d/%m/%Y')} - {to_date.strftime('%d/%m/%Y')}"
            next_row = cls._add_title_section(
                ws,
                "Delivery Invoices Report",
                f"Date Range: {date_range} | Generated: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            )

            # Headers
            headers = [
                "Invoice Number", "Date", "Rice Mill", "Mill Code",
                "Truck Number", "Bills Count", "Total Weight (kg)", "Created By"
            ]

            for col, header in enumerate(headers, 1):
                ws.cell(row=next_row, column=col, value=header)

            cls._format_header_row(ws, next_row, len(headers))

            # Data rows
            data_start_row = next_row + 1
            for row_idx, invoice in enumerate(invoices, data_start_row):
                # Get bills count for this invoice
                bills = db.query(PurchaseBill).filter(
                    PurchaseBill.delivery_invoice_id == invoice.id
                ).all()

                ws.cell(row=row_idx, column=1, value=invoice.invoice_number)
                ws.cell(row=row_idx, column=2, value=invoice.invoice_date)
                ws.cell(row=row_idx, column=3, value=invoice.mill.mill_name if invoice.mill else "Unknown")
                ws.cell(row=row_idx, column=4, value=invoice.mill.mill_code if invoice.mill and invoice.mill.mill_code else "N/A")
                ws.cell(row=row_idx, column=5, value=invoice.truck.truck_number if invoice.truck else "Unknown")
                ws.cell(row=row_idx, column=6, value=len(bills))
                ws.cell(row=row_idx, column=7, value=float(invoice.total_weight))
                ws.cell(row=row_idx, column=8, value=invoice.created_by or "System")

            # Format columns
            data_end_row = data_start_row + len(invoices) - 1
            if len(invoices) > 0:
                cls._format_date_column(ws, 2, data_start_row, data_end_row)  # Date
                cls._format_number_column(ws, 7, data_start_row, data_end_row)  # Total Weight

            # Add summary row
            summary_row = data_end_row + 2
            ws.cell(row=summary_row, column=1, value="TOTAL")
            ws.cell(row=summary_row, column=1).font = Font(bold=True)

            total_weight = sum(float(inv.total_weight) for inv in invoices)
            total_bills = sum(len(db.query(PurchaseBill).filter(
                PurchaseBill.delivery_invoice_id == inv.id
            ).all()) for inv in invoices)

            ws.cell(row=summary_row, column=6, value=total_bills)
            ws.cell(row=summary_row, column=7, value=total_weight)

            ws.cell(row=summary_row, column=6).font = Font(bold=True)
            ws.cell(row=summary_row, column=7).font = Font(bold=True)
            ws.cell(row=summary_row, column=7).number_format = cls.NUMBER_FORMAT

            # Auto-adjust column widths
            cls._auto_adjust_column_width(ws)

            # Save file
            wb.save(file_path)
            return file_path

        except Exception as e:
            raise ExcelExportError(f"Failed to export delivery invoices: {str(e)}")

    @classmethod
    def export_farmer_summary(
        cls,
        db: Session,
        file_path: str,
        from_date: datetime,
        to_date: datetime,
        filter_id: Optional[int] = None
    ) -> str:
        """
        Export farmer summary report to Excel

        Args:
            db: Database session
            file_path: Output file path
            from_date: Start date for filtering
            to_date: End date for filtering
            filter_id: Optional farmer ID to filter by

        Returns:
            Path to the generated Excel file
        """
        try:
            # Build query
            query = db.query(
                PurchaseBill.farmer_id,
                func.count(PurchaseBill.id).label('bill_count'),
                func.sum(PurchaseBill.net_weight).label('total_weight'),
                func.sum(PurchaseBill.total_payment).label('total_payment')
            ).filter(
                PurchaseBill.bill_date >= from_date,
                PurchaseBill.bill_date <= to_date
            )

            if filter_id:
                query = query.filter(PurchaseBill.farmer_id == filter_id)

            farmer_stats = query.group_by(PurchaseBill.farmer_id).all()

            # Create workbook
            wb = cls._create_workbook()
            ws = wb.active
            ws.title = "Farmer Summary"

            # Add title section
            date_range = f"{from_date.strftime('%d/%m/%Y')} - {to_date.strftime('%d/%m/%Y')}"
            next_row = cls._add_title_section(
                ws,
                "Farmer Summary Report",
                f"Date Range: {date_range} | Generated: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            )

            # Headers
            headers = [
                "Farmer IC", "Farmer Name", "Phone", "Address",
                "Bills Count", "Total Weight (kg)", "Total Payment (RM)", "Avg Weight per Bill (kg)"
            ]

            for col, header in enumerate(headers, 1):
                ws.cell(row=next_row, column=col, value=header)

            cls._format_header_row(ws, next_row, len(headers))

            # Data rows
            data_start_row = next_row + 1
            for row_idx, stat in enumerate(farmer_stats, data_start_row):
                farmer = db.query(Farmer).filter(Farmer.id == stat.farmer_id).first()

                if not farmer:
                    continue

                total_weight = float(stat.total_weight) if stat.total_weight else 0
                total_payment = float(stat.total_payment) if stat.total_payment else 0
                avg_weight = total_weight / stat.bill_count if stat.bill_count > 0 else 0

                ws.cell(row=row_idx, column=1, value=farmer.ic_number)
                ws.cell(row=row_idx, column=2, value=farmer.name)
                ws.cell(row=row_idx, column=3, value=farmer.phone_number or "N/A")
                ws.cell(row=row_idx, column=4, value=farmer.address or "N/A")
                ws.cell(row=row_idx, column=5, value=stat.bill_count)
                ws.cell(row=row_idx, column=6, value=total_weight)
                ws.cell(row=row_idx, column=7, value=total_payment)
                ws.cell(row=row_idx, column=8, value=avg_weight)

            # Format columns
            data_end_row = data_start_row + len(farmer_stats) - 1
            if len(farmer_stats) > 0:
                cls._format_number_column(ws, 6, data_start_row, data_end_row)  # Total Weight
                cls._format_currency_column(ws, 7, data_start_row, data_end_row)  # Total Payment
                cls._format_number_column(ws, 8, data_start_row, data_end_row)  # Avg Weight

            # Add summary row
            summary_row = data_end_row + 2
            ws.cell(row=summary_row, column=1, value="GRAND TOTAL")
            ws.cell(row=summary_row, column=1).font = Font(bold=True)

            total_bills = sum(s.bill_count for s in farmer_stats)
            total_weight = sum(float(s.total_weight) if s.total_weight else 0 for s in farmer_stats)
            total_payment = sum(float(s.total_payment) if s.total_payment else 0 for s in farmer_stats)

            ws.cell(row=summary_row, column=5, value=total_bills)
            ws.cell(row=summary_row, column=6, value=total_weight)
            ws.cell(row=summary_row, column=7, value=total_payment)

            ws.cell(row=summary_row, column=5).font = Font(bold=True)
            ws.cell(row=summary_row, column=6).font = Font(bold=True)
            ws.cell(row=summary_row, column=7).font = Font(bold=True)

            ws.cell(row=summary_row, column=6).number_format = cls.NUMBER_FORMAT
            ws.cell(row=summary_row, column=7).number_format = cls.CURRENCY_FORMAT

            # Auto-adjust column widths
            cls._auto_adjust_column_width(ws)

            # Save file
            wb.save(file_path)
            return file_path

        except Exception as e:
            raise ExcelExportError(f"Failed to export farmer summary: {str(e)}")

    @classmethod
    def export_mill_summary(
        cls,
        db: Session,
        file_path: str,
        from_date: datetime,
        to_date: datetime,
        filter_id: Optional[int] = None
    ) -> str:
        """
        Export mill summary report to Excel

        Args:
            db: Database session
            file_path: Output file path
            from_date: Start date for filtering
            to_date: End date for filtering
            filter_id: Optional mill ID to filter by

        Returns:
            Path to the generated Excel file
        """
        try:
            # Build query
            query = db.query(
                DeliveryInvoice.mill_id,
                func.count(DeliveryInvoice.id).label('invoice_count'),
                func.sum(DeliveryInvoice.total_weight).label('total_weight')
            ).filter(
                DeliveryInvoice.invoice_date >= from_date,
                DeliveryInvoice.invoice_date <= to_date
            )

            if filter_id:
                query = query.filter(DeliveryInvoice.mill_id == filter_id)

            mill_stats = query.group_by(DeliveryInvoice.mill_id).all()

            # Create workbook
            wb = cls._create_workbook()
            ws = wb.active
            ws.title = "Mill Summary"

            # Add title section
            date_range = f"{from_date.strftime('%d/%m/%Y')} - {to_date.strftime('%d/%m/%Y')}"
            next_row = cls._add_title_section(
                ws,
                "Rice Mill Summary Report",
                f"Date Range: {date_range} | Generated: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            )

            # Headers
            headers = [
                "Mill Code", "Mill Name", "Location", "Contact",
                "Invoices Count", "Total Weight (kg)", "Avg Weight per Invoice (kg)"
            ]

            for col, header in enumerate(headers, 1):
                ws.cell(row=next_row, column=col, value=header)

            cls._format_header_row(ws, next_row, len(headers))

            # Data rows
            data_start_row = next_row + 1
            for row_idx, stat in enumerate(mill_stats, data_start_row):
                mill = db.query(RiceMill).filter(RiceMill.id == stat.mill_id).first()

                if not mill:
                    continue

                total_weight = float(stat.total_weight) if stat.total_weight else 0
                avg_weight = total_weight / stat.invoice_count if stat.invoice_count > 0 else 0

                ws.cell(row=row_idx, column=1, value=mill.mill_code or "N/A")
                ws.cell(row=row_idx, column=2, value=mill.mill_name)
                ws.cell(row=row_idx, column=3, value=mill.location or "N/A")
                ws.cell(row=row_idx, column=4, value=mill.contact_number or "N/A")
                ws.cell(row=row_idx, column=5, value=stat.invoice_count)
                ws.cell(row=row_idx, column=6, value=total_weight)
                ws.cell(row=row_idx, column=7, value=avg_weight)

            # Format columns
            data_end_row = data_start_row + len(mill_stats) - 1
            if len(mill_stats) > 0:
                cls._format_number_column(ws, 6, data_start_row, data_end_row)  # Total Weight
                cls._format_number_column(ws, 7, data_start_row, data_end_row)  # Avg Weight

            # Add summary row
            summary_row = data_end_row + 2
            ws.cell(row=summary_row, column=1, value="GRAND TOTAL")
            ws.cell(row=summary_row, column=1).font = Font(bold=True)

            total_invoices = sum(s.invoice_count for s in mill_stats)
            total_weight = sum(float(s.total_weight) if s.total_weight else 0 for s in mill_stats)

            ws.cell(row=summary_row, column=5, value=total_invoices)
            ws.cell(row=summary_row, column=6, value=total_weight)

            ws.cell(row=summary_row, column=5).font = Font(bold=True)
            ws.cell(row=summary_row, column=6).font = Font(bold=True)
            ws.cell(row=summary_row, column=6).number_format = cls.NUMBER_FORMAT

            # Auto-adjust column widths
            cls._auto_adjust_column_width(ws)

            # Save file
            wb.save(file_path)
            return file_path

        except Exception as e:
            raise ExcelExportError(f"Failed to export mill summary: {str(e)}")
